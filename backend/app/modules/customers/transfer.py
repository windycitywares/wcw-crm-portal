"""Customer export (CSV / Excel) and import (CSV / Excel), for managers and admins (`data:export`).

Export: the list's filters apply; one row per customer, custom fields as `cf_<key>` columns.
Import: columns are matched by name (see COLUMNS, case and spaces ignored). Rows are matched to existing customers
by email. Mode "skip" leaves existing customers alone; mode "update" fills in the values the file gives.
A dry run checks everything and changes nothing. Bad rows are reported and skipped; the rest is imported.
Custom fields, consent dates and shop numbers are export-only (consent can't be granted by a file).
"""

from __future__ import annotations

import csv
import io
from datetime import datetime
from decimal import Decimal
from typing import Any, Literal

from pydantic import BaseModel, EmailStr, TypeAdapter, ValidationError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import ConflictError, NotFoundError, ValidationFailedError
from app.modules.audit.service import AuditService
from app.modules.customfields.service import CustomFieldService
from app.modules.lookups.service import LookupService
from app.modules.tags.service import TagService
from app.modules.users.service import UserService

from .repository import CustomerRepository
from .service import ENTITY, CustomerService

MAX_EXPORT = 50_000
MAX_IMPORT_ROWS = 5_000
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_ERRORS = 100

EXPORT_COLUMNS = (
    "id", "first_name", "last_name", "company", "email", "other_emails", "phone", "type", "source", "status", "owner",
    "marketing_consent", "consent_at", "unsubscribed_at", "tags", "address_line1", "address_line2", "city", "state",
    "postal_code", "country", "orders_count", "total_spent", "first_order_at", "last_order_at", "shop_user_id",
    "created_at", "updated_at",
)

# Import column → accepted header spellings (lowercase, spaces/underscores/dashes removed).
COLUMNS: dict[str, tuple[str, ...]] = {
    "first_name": ("firstname", "first", "givenname"),
    "last_name": ("lastname", "last", "surname", "familyname"),
    "full_name": ("name", "fullname", "customername", "contactname"),
    "company": ("company", "companyname", "business", "organization", "organisation"),
    "email": ("email", "emailaddress", "mail"),
    "other_emails": ("otheremails", "altemails", "alternateemails"),
    "phone": ("phone", "phonenumber", "telephone", "mobile", "tel"),
    "type": ("type", "customertype"),
    "source": ("source", "leadsource"),
    "status": ("status",),
    "owner": ("owner", "owneremail", "accountowner"),
    "tags": ("tags", "tag", "labels"),
    "address_line1": ("addressline1", "address1", "address", "street", "line1"),
    "address_line2": ("addressline2", "address2", "line2"),
    "city": ("city", "town"),
    "state": ("state", "region", "province", "county"),
    "postal_code": ("postalcode", "zip", "zipcode", "postcode"),
    "country": ("country",),
}
ADDRESS = {"address_line1": "line1", "address_line2": "line2", "city": "city", "state": "state",
           "postal_code": "postal_code", "country": "country"}
STATUSES = ("active", "inactive", "blocked")
_email = TypeAdapter(EmailStr)


def _norm(header: str) -> str:
    return "".join(ch for ch in header.strip().lower() if ch.isalnum())


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    return str(value)


def _safe(text: str) -> str:
    """Stop spreadsheet apps running a cell as a formula (CSV injection)."""
    return f"'{text}" if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


class ImportResult(BaseModel):
    dry_run: bool
    rows: int
    created: int
    updated: int
    skipped: int
    errors: list[dict[str, Any]]
    columns: dict[str, str]  # file header → CRM field
    ignored_columns: list[str]


class CustomerTransfer:
    def __init__(self, session: AsyncSession):
        self.session = session
        self.repo = CustomerRepository(session)
        self.customers = CustomerService(session)
        self.lookups = LookupService(session)
        self.tags = TagService(session)
        self.fields = CustomFieldService(session)
        self.audit = AuditService(session)

    # ---- export ---------------------------------------------------------------------------
    async def _owners(self) -> tuple[dict[int, str], dict[str, int]]:
        people = await UserService(self.session).contacts()
        by_id = {uid: email for uid, _, email in people}
        by_key: dict[str, int] = {}
        for uid, name, email in people:
            by_key.setdefault(name.strip().lower(), uid)
            by_key[email.strip().lower()] = uid
        return by_id, by_key

    async def export_rows(self, filters: dict[str, Any], actor) -> tuple[list[str], list[list[str]]]:
        tag_id = filters.pop("tag_id", None)
        ids = await self.tags.entity_ids_with_tag(ENTITY, tag_id) if tag_id else None
        rows = list(await self.repo.list(self.repo.search(ids=ids, **filters).limit(MAX_EXPORT + 1)))
        if len(rows) > MAX_EXPORT:
            raise ValidationFailedError(f"More than {MAX_EXPORT:,} customers match: narrow the filters and export in parts")
        tags = await self.tags.tags_for(ENTITY, [c.id for c in rows])
        owners, _ = await self._owners()
        types = {v.key: v.label for v in await self.lookups.values("customer_type")}
        sources = {v.key: v.label for v in await self.lookups.values("lead_source")}
        defs = [d for d in await self.fields.repo.for_entity(ENTITY)]
        header = [*EXPORT_COLUMNS, *(f"cf_{d.key}" for d in defs)]
        out: list[list[str]] = []
        for c in rows:
            a = c.address or {}
            values = {
                "id": c.id, "first_name": c.first_name, "last_name": c.last_name, "company": c.company, "email": c.email,
                "other_emails": c.alt_emails, "phone": c.phone, "type": types.get(c.customer_type or "", c.customer_type),
                "source": sources.get(c.source or "", c.source), "status": c.status,
                "owner": owners.get(c.owner_id) if c.owner_id else None, "marketing_consent": c.marketing_consent,
                "consent_at": c.consent_at, "unsubscribed_at": c.unsubscribed_at,
                "tags": [t.name for t in tags.get(c.id, [])],
                **{col: a.get(key) for col, key in ADDRESS.items()},
                "orders_count": c.orders_count, "total_spent": c.total_spent, "first_order_at": c.first_order_at,
                "last_order_at": c.last_order_at, "shop_user_id": c.shop_user_id, "created_at": c.created_at,
                "updated_at": c.updated_at,
            }
            cf = c.custom_fields or {}
            out.append([_safe(_cell(values[k])) for k in EXPORT_COLUMNS] + [_safe(_cell(cf.get(d.key))) for d in defs])
        await self.audit.record("customer.export", actor=actor, entity_type=ENTITY, changes={"rows": len(out)})
        await self.session.commit()
        return header, out

    @staticmethod
    def to_csv(header: list[str], rows: list[list[str]]) -> bytes:
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(header)
        writer.writerows(rows)
        return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM: Excel opens UTF-8 correctly

    @staticmethod
    def to_xlsx(header: list[str], rows: list[list[str]]) -> bytes:
        from openpyxl import Workbook

        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Customers")
        ws.append(header)
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ---- import ---------------------------------------------------------------------------
    @staticmethod
    def read_table(filename: str, data: bytes) -> list[list[str]]:
        if len(data) > MAX_IMPORT_BYTES:
            raise ValidationFailedError("The file is larger than 5 MB: split it into smaller files")
        if filename.lower().endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook

            try:
                wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            except Exception:  # noqa: BLE001 - any parse failure is the same answer for the user
                raise ValidationFailedError("Couldn't read this Excel file") from None
            ws = wb.worksheets[0]
            return [["" if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)]
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValidationFailedError("Couldn't read the file: save it as CSV (UTF-8) or Excel")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]

    @staticmethod
    def map_columns(header: list[str]) -> tuple[dict[int, str], list[str]]:
        lookup = {alias: field for field, aliases in COLUMNS.items() for alias in (*aliases, _norm(field))}
        mapping: dict[int, str] = {}
        ignored: list[str] = []
        for i, h in enumerate(header):
            field = lookup.get(_norm(h))
            if field and field not in mapping.values():
                mapping[i] = field
            elif h:
                ignored.append(h)
        return mapping, ignored

    @staticmethod
    def _choice(value: str, options: dict[str, str], what: str) -> str:
        key = options.get(value.strip().lower())
        if key is None:
            raise ValidationFailedError(f"Unknown {what} '{value}' (add it in Administration → Types first)")
        return key

    async def import_file(self, filename: str, data: bytes, actor, *, mode: Literal["skip", "update"] = "skip",
                          dry_run: bool = True) -> ImportResult:
        table = [r for r in self.read_table(filename, data) if any(c for c in r)]
        if not table:
            raise ValidationFailedError("The file is empty")
        header, body = table[0], table[1:]
        if len(body) > MAX_IMPORT_ROWS:
            raise ValidationFailedError(f"At most {MAX_IMPORT_ROWS:,} rows per file: split it")
        mapping, ignored = self.map_columns(header)
        fields = set(mapping.values())
        if not fields & {"email", "first_name", "last_name", "full_name", "company", "phone"}:
            raise ValidationFailedError("No usable columns: the file needs at least an Email or Name column")

        def options(kind_values) -> dict[str, str]:
            out: dict[str, str] = {}
            for v in kind_values:
                if v.is_active:
                    out[v.key.lower()] = v.key
                    out[v.label.strip().lower()] = v.key
            return out

        types = options(await self.lookups.values("customer_type"))
        sources = options(await self.lookups.values("lead_source"))
        _, owners = await self._owners()
        result = ImportResult(dry_run=dry_run, rows=len(body), created=0, updated=0, skipped=0, errors=[],
                              columns={header[i]: f for i, f in mapping.items()}, ignored_columns=ignored)
        seen_emails: set[str] = set()

        for n, raw in enumerate(body, start=2):  # row 1 is the header
            row = {f: (raw[i] if i < len(raw) else "").strip() for i, f in mapping.items()}
            try:
                values, tag_names = self._row_values(row, types, sources, owners)
                email = values.get("email")
                if email:
                    if email in seen_emails:
                        raise ValidationFailedError("This email appears twice in the file")
                    seen_emails.add(email)
                async with self.session.begin_nested():
                    outcome = await self.customers.import_row(values, actor, update_existing=mode == "update")
                    if outcome[0] != "skipped" and tag_names:
                        have = [t.name for t in (await self.tags.tags_for(ENTITY, [outcome[1]])).get(outcome[1], [])]
                        await self.tags.set_for(actor, ENTITY, outcome[1], have + tag_names, commit=False)
                setattr(result, outcome[0], getattr(result, outcome[0]) + 1)
            except (ValidationFailedError, ConflictError, NotFoundError, ValidationError) as exc:
                result.skipped += 1
                if len(result.errors) < MAX_ERRORS:
                    message = exc.errors()[0]["msg"] if isinstance(exc, ValidationError) else str(exc)
                    result.errors.append({"row": n, "error": message.removeprefix("Value error, ")})

        if dry_run:
            await self.session.rollback()
        else:
            await self.audit.record("customer.import", actor=actor, entity_type=ENTITY, changes={
                "file": filename[:100], "created": result.created, "updated": result.updated, "skipped": result.skipped})
            await self.session.commit()
        return result

    def _row_values(self, row: dict[str, str], types: dict[str, str], sources: dict[str, str],
                    owners: dict[str, int]) -> tuple[dict[str, Any], list[str]]:
        values: dict[str, Any] = {}
        if row.get("first_name") or row.get("last_name"):
            values["first_name"], values["last_name"] = row.get("first_name", ""), row.get("last_name", "")
        elif row.get("full_name"):
            first, _, last = " ".join(row["full_name"].split()).partition(" ")
            values["first_name"], values["last_name"] = first, last
        for f in ("company", "phone"):
            if row.get(f):
                values[f] = row[f]
        if row.get("email"):
            values["email"] = str(_email.validate_python(row["email"])).lower()
        if row.get("other_emails"):
            values["alt_emails"] = [str(_email.validate_python(e)).lower()
                                    for e in row["other_emails"].replace(";", ",").split(",") if e.strip()]
        if row.get("type"):
            values["customer_type"] = self._choice(row["type"], types, "customer type")
        if row.get("source"):
            values["source"] = self._choice(row["source"], sources, "source")
        if row.get("status"):
            status = row["status"].lower()
            if status not in STATUSES:
                raise ValidationFailedError(f"Status must be one of: {', '.join(STATUSES)}")
            values["status"] = status
        if row.get("owner"):
            owner = owners.get(row["owner"].lower())
            if owner is None:
                raise ValidationFailedError(f"No colleague called '{row['owner']}' (use their email or full name)")
            values["owner_id"] = owner
        address = {key: row[col] for col, key in ADDRESS.items() if row.get(col)}
        if address:
            values["address"] = address
        for f, limit in (("first_name", 80), ("last_name", 80), ("company", 160), ("phone", 40)):
            if len(values.get(f) or "") > limit:
                raise ValidationFailedError(f"{f.replace('_', ' ').capitalize()} is longer than {limit} characters")
        if not any(values.get(k) for k in ("first_name", "last_name", "company", "email", "phone")):
            raise ValidationFailedError("The row has no name, company, email or phone")
        tag_names = [t.strip() for t in row.get("tags", "").replace(";", ",").split(",") if t.strip()][:20]
        return values, tag_names
